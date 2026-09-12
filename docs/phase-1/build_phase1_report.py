from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "Bao_cao_giai_doan_1_phan_tich_thiet_ke_he_thong.docx"

BLACK = "000000"
DARK = "1F2933"
GRAY = "4B5563"
LIGHT_GRAY = "F3F4F6"
BORDER = "D9D9D9"
BLUE = "1F4E79"
PALE_BLUE = "EAF3F8"


def set_font(run, size=10.5, bold=False, italic=False, color=DARK, name="Aptos"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=120, start=130, bottom=120, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_cell_text(cell, text, bold=False, fill=None, color=DARK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)
    if fill:
        shade(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    set_font(r, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths=None, header_fill=BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0]
    tr_pr = hdr._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for i, header in enumerate(headers):
        set_cell_text(hdr.cells[i], header, bold=True, fill=header_fill, color="FFFFFF", size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            hdr.cells[i].width = Inches(widths[i])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        fill = "FFFFFF" if row_index % 2 == 0 else PALE_BLUE
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) <= 18 and "\n" not in str(value) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], value, fill=fill, size=8.8, align=align)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(item)
        set_font(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(item)
        set_font(r)


def add_manual_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(f"{index}.  {item}")
        set_font(r)


def para(doc, text="", bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead:
        lead = p.add_run(bold_lead)
        set_font(lead, bold=True)
        body = p.add_run(text)
        set_font(body)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    remove_paragraph_borders(p)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        set_font(run, size=15 if level == 1 else 12 if level == 2 else 10.5, bold=True, color=BLACK, name="Aptos Display")
    return p


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def remove_style_borders(style):
    p_pr = style._element.find(qn("w:pPr"))
    if p_pr is None:
        return
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def page_break(doc):
    doc.add_page_break()


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_font(run, size=8.5, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        s = styles[style_name]
        remove_style_borders(s)
        s.font.name = "Aptos Display"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        s.font.color.rgb = RGBColor.from_string(BLACK)

    for style_name in ["List Bullet", "List Number"]:
        s = styles[style_name]
        s.font.name = "Aptos"
        s.font.size = Pt(10.5)
        s.paragraph_format.left_indent = Inches(0.28)
        s.paragraph_format.first_line_indent = Inches(-0.14)

    header = section.header.paragraphs[0]
    header.text = "Báo cáo Giai đoạn 1 | Website thương mại điện tử mỹ phẩm"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_font(run, size=8.2, color=GRAY)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(52)
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("Báo cáo Giai đoạn 1")
    set_font(r, size=22, bold=True, color=BLACK, name="Aptos Display")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    r = subtitle.add_run("Phân tích yêu cầu, thiết kế hệ thống và chuẩn bị dữ liệu cho website thương mại điện tử mỹ phẩm tích hợp hệ thống gợi ý, chatbot và analytics")
    set_font(r, size=13, bold=True, color=DARK, name="Aptos Display")

    info_rows = [
        ("Tên đề tài", "Xây dựng website thương mại điện tử mỹ phẩm tích hợp hệ thống gợi ý lai cá nhân hóa và chatbot tư vấn dựa trên dữ liệu hệ thống"),
        ("Loại báo cáo", "Báo cáo triển khai Giai đoạn 1"),
        ("Nhóm thực hiện", "Nhóm 3"),
        ("Thành viên", "Trần Văn Nhật - 23010625\nNguyễn Huy Hiệp - 23010178"),
        ("Giảng viên hướng dẫn", "Đặng Thị Thúy An"),
        ("Ngày lập", "10/09/2026"),
    ]
    add_table(doc, ["Thông tin", "Nội dung"], info_rows, widths=[1.6, 5.4], header_fill="404040")

    para(doc, "Báo cáo này tổng hợp các việc đã thực hiện trong Giai đoạn 1 của đồ án. Trọng tâm của giai đoạn này là thống nhất phạm vi hệ thống, chia tách source để làm việc nhóm, xác định kiến trúc tổng thể, thiết kế use case, thiết kế cơ sở dữ liệu ban đầu, định nghĩa API liên kết giữa website và các service hệ thống, đồng thời chuẩn bị khung dữ liệu sản phẩm mỹ phẩm phục vụ recommendation, chatbot và analytics.")

    page_break(doc)

    heading(doc, "Mục lục", 1)
    add_manual_numbered(doc, [
        "Mục tiêu và phạm vi Giai đoạn 1",
        "Phân tích yêu cầu chức năng và phi chức năng",
        "Đối tượng sử dụng hệ thống",
        "Phân chia source và trách nhiệm phát triển",
        "Kiến trúc tổng thể hệ thống",
        "Use Case chính",
        "Thiết kế cơ sở dữ liệu ban đầu",
        "Thiết kế API và event contract",
        "Chuẩn bị Product Dataset",
        "Kết quả đã hoàn thành và kế hoạch tiếp theo",
    ])

    heading(doc, "1 Mục tiêu và phạm vi Giai đoạn 1", 1)
    para(doc, "Giai đoạn 1 đóng vai trò nền móng cho toàn bộ quá trình phát triển. Mục tiêu không phải là hoàn thiện ngay website hoặc mô hình AI, mà là tạo ra một bản thiết kế đủ rõ để các thành viên có thể làm song song, giảm phụ thuộc lẫn nhau và hạn chế sửa lại kiến trúc khi bước sang giai đoạn code.")
    add_table(doc, ["Mục tiêu", "Kết quả cần có"], [
        ("Phân tích nghiệp vụ", "Xác định chức năng cho khách hàng, quản trị viên và các service hệ thống."),
        ("Tách source", "Tổ chức repo theo monorepo, phân chia rõ khu vực website và khu vực hệ thống bên trong."),
        ("Thiết kế kiến trúc", "Xác định các khối chính gồm web app, backend/API gateway, recommendation, chatbot, analytics, stream analytics và database."),
        ("Thiết kế dữ liệu", "Tạo schema database ban đầu, data dictionary và schema dataset sản phẩm mỹ phẩm."),
        ("Chuẩn hóa tích hợp", "Tạo OpenAPI contract và behavior event contract để website kết nối với hệ thống."),
    ], widths=[2.0, 4.9])

    heading(doc, "2 Phân tích yêu cầu chức năng và phi chức năng", 1)
    heading(doc, "2.1 Yêu cầu chức năng", 2)
    add_table(doc, ["Nhóm", "Chức năng cụ thể", "Ý nghĩa triển khai"], [
        ("Khách hàng", "Đăng ký, đăng nhập, quản lý tài khoản", "Là nền tảng để cá nhân hóa trải nghiệm và lưu lịch sử mua hàng."),
        ("Khách hàng", "Xem danh mục, chi tiết sản phẩm, tìm kiếm và lọc", "Giúp người dùng tiếp cận sản phẩm theo thương hiệu, danh mục, giá, loại da và vấn đề da."),
        ("Khách hàng", "Yêu thích, giỏ hàng, đặt hàng và thanh toán mô phỏng", "Hoàn thiện luồng thương mại điện tử cốt lõi trong phạm vi đồ án cơ sở."),
        ("Khách hàng", "Đánh giá sản phẩm", "Tạo dữ liệu phản hồi cho sản phẩm, analytics và collaborative filtering."),
        ("Khách hàng", "Nhận gợi ý cá nhân hóa", "Hiển thị sản phẩm phù hợp dựa trên hồ sơ nhu cầu và hành vi."),
        ("Khách hàng", "Chat với chatbot tư vấn", "Hỗ trợ hỏi đáp sản phẩm dựa trên dữ liệu đã lưu trong hệ thống."),
        ("Admin", "Quản lý sản phẩm, danh mục, thương hiệu", "Đảm bảo dữ liệu sản phẩm có thể được cập nhật và kiểm duyệt."),
        ("Admin", "Quản lý đơn hàng, người dùng, đánh giá", "Hỗ trợ vận hành website và theo dõi hoạt động mua hàng."),
        ("Admin", "Xem dashboard phân tích", "Theo dõi doanh thu, hành vi người dùng, hiệu quả gợi ý và chatbot."),
        ("Hệ thống", "Ghi nhận behavior event", "Lưu lại hành vi để phục vụ recommendation, analytics và stream analytics."),
        ("Hệ thống", "Cập nhật user profile", "Tổng hợp loại da, nhu cầu, sở thích, ngân sách và hành vi quan tâm."),
        ("Hệ thống", "Tính điểm gợi ý", "Kết hợp Knowledge-Based, Content-Based và Item-Based Collaborative Filtering."),
    ], widths=[1.2, 2.45, 3.25])

    heading(doc, "2.2 Yêu cầu phi chức năng", 2)
    add_table(doc, ["Nhóm yêu cầu", "Mô tả"], [
        ("Hiệu năng", "Các API danh sách sản phẩm, chi tiết sản phẩm, gợi ý và dashboard cần đủ nhanh cho môi trường demo local. Dữ liệu hành vi cần có index theo user, product, event type và thời gian."),
        ("Bảo mật", "Không lưu mật khẩu dạng thô. Hệ thống cần phân quyền CUSTOMER và ADMIN. File cấu hình local không chứa mật khẩu thật khi đưa lên repository."),
        ("Toàn vẹn dữ liệu", "Sản phẩm cần có nguồn dữ liệu, ngày xác minh, danh mục, thương hiệu và các thuộc tính quan trọng cho hệ thống gợi ý."),
        ("Khả mở rộng", "Source được tách thành website, recommendation, chatbot, analytics, stream analytics, database và dataset để có thể mở rộng từng phần."),
        ("Khả bảo trì", "API contract, event contract và migration database được lưu riêng để nhóm có thể kiểm soát thay đổi."),
        ("An toàn tư vấn", "Chatbot không chẩn đoán bệnh da, không cam kết điều trị và không bịa thông tin ngoài dữ liệu đã kiểm duyệt."),
    ], widths=[1.55, 5.35])

    heading(doc, "3 Đối tượng sử dụng hệ thống", 1)
    add_table(doc, ["Đối tượng", "Nhu cầu", "Tính năng liên quan"], [
        ("Khách hàng mới", "Chưa có lịch sử hành vi, cần tìm sản phẩm phù hợp nhanh.", "Quiz nhu cầu, bộ lọc, Knowledge-Based Recommendation, chatbot."),
        ("Khách hàng quay lại", "Muốn nhận gợi ý dựa trên sở thích, hành vi xem, yêu thích, giỏ hàng và mua hàng.", "User profile, Content-Based, Collaborative Filtering, gợi ý cá nhân hóa."),
        ("Quản trị viên", "Cần quản lý dữ liệu và theo dõi hiệu quả hoạt động website.", "Admin dashboard, quản lý sản phẩm, đơn hàng, người dùng, đánh giá."),
        ("Hệ thống recommendation", "Cần dữ liệu sản phẩm và hành vi để tính điểm phù hợp.", "Product dataset, behavior events, user profiles, recommendation logs."),
        ("Hệ thống chatbot", "Cần dữ liệu sản phẩm có cấu trúc để tư vấn chính xác.", "Product attributes, ingredient data, warnings, recommendation API."),
        ("Hệ thống analytics", "Cần dữ liệu giao dịch và hành vi để tạo dashboard.", "Orders, order items, behavior events, recommendation logs, chatbot conversations."),
    ], widths=[1.45, 2.75, 2.7])

    heading(doc, "4 Phân chia source và trách nhiệm phát triển", 1)
    para(doc, "Dự án được clone từ repository GitHub và triển khai trong thư mục DACS trên máy local. Repository được tổ chức theo hướng monorepo để mỗi thành viên có khu vực làm việc riêng nhưng vẫn có điểm liên kết rõ ràng.")
    add_table(doc, ["Khu vực source", "Người phụ trách chính", "Nội dung"], [
        ("apps/web", "Thành viên làm website", "Giao diện khách hàng, giao diện admin, routing, gọi API và gửi behavior event."),
        ("services/recommendation-service", "Thành viên làm hệ thống", "Knowledge-Based, Content-Based, Collaborative Filtering, Hybrid Scoring và giải thích lý do gợi ý."),
        ("services/chatbot-service", "Thành viên làm hệ thống", "Chatbot tư vấn sản phẩm dựa trên dữ liệu sản phẩm, luật an toàn và kết quả recommendation."),
        ("services/analytics-service", "Thành viên làm hệ thống", "Tổng hợp dữ liệu cho dashboard và đánh giá hiệu quả recommendation/chatbot."),
        ("services/stream-analytics-service", "Thành viên làm hệ thống", "Nền tảng xử lý event gần thời gian thực, dùng để mở rộng khi hệ thống ổn định."),
        ("database", "Thành viên làm hệ thống", "Migration SQL, seed data, data dictionary và quy tắc database."),
        ("datasets", "Thành viên làm hệ thống", "Dữ liệu sản phẩm thô, dữ liệu đã làm sạch và schema product catalog."),
        ("contracts", "Cả nhóm thống nhất", "OpenAPI và event contract để website và hệ thống tích hợp đúng."),
        ("docs", "Cả nhóm", "Tài liệu kiến trúc, phân tích, thiết kế và kế hoạch triển khai."),
    ], widths=[2.0, 1.6, 3.3])

    heading(doc, "5 Kiến trúc tổng thể hệ thống", 1)
    para(doc, "Kiến trúc được chia thành các lớp độc lập. Website là lớp giao diện. Backend hoặc API gateway là lớp tiếp nhận request và điều phối. Các service hệ thống xử lý recommendation, chatbot và analytics. Database và dataset là lớp dữ liệu trung tâm.")
    add_table(doc, ["Lớp", "Thành phần", "Vai trò"], [
        ("Presentation", "apps/web", "Hiển thị website, admin dashboard, chatbot UI và các khối gợi ý sản phẩm."),
        ("Integration", "Backend API hoặc API Gateway", "Nhận request từ website, kiểm tra quyền, gọi module/service phù hợp và trả response."),
        ("Business", "Ecommerce modules", "Xử lý sản phẩm, danh mục, tài khoản, giỏ hàng, đơn hàng và đánh giá."),
        ("AI System", "recommendation-service", "Tính sản phẩm phù hợp theo KB, CB, CF và Hybrid Scoring."),
        ("AI System", "chatbot-service", "Tư vấn sản phẩm bằng dữ liệu đã kiểm duyệt và các mẫu phản hồi an toàn."),
        ("Analytics", "analytics-service", "Tổng hợp số liệu dashboard, conversion, top sản phẩm và hiệu quả gợi ý."),
        ("Streaming", "stream-analytics-service", "Xử lý behavior event gần thời gian thực khi cần mở rộng."),
        ("Data", "MySQL, datasets", "Lưu dữ liệu vận hành, dữ liệu hành vi, log AI và product dataset."),
    ], widths=[1.35, 2.05, 3.5])

    heading(doc, "6 Use Case chính", 1)
    add_table(doc, ["Mã", "Actor", "Use Case", "Kết quả mong đợi"], [
        ("UC01", "Khách hàng", "Đăng ký và đăng nhập", "Người dùng có tài khoản hợp lệ và có thể sử dụng chức năng cá nhân hóa."),
        ("UC02", "Khách hàng", "Tìm kiếm và lọc sản phẩm", "Danh sách sản phẩm được lọc theo nhu cầu, giá, thương hiệu, loại da và vấn đề da."),
        ("UC03", "Khách hàng", "Xem chi tiết sản phẩm", "Hiển thị thông tin sản phẩm, thành phần, công dụng, cảnh báo và sản phẩm tương tự."),
        ("UC04", "Khách hàng", "Nhận gợi ý cá nhân hóa", "Hệ thống trả về danh sách sản phẩm kèm điểm và lý do gợi ý."),
        ("UC05", "Khách hàng", "Chatbot tư vấn", "Chatbot phản hồi dựa trên dữ liệu sản phẩm và không vượt ngoài phạm vi tư vấn an toàn."),
        ("UC06", "Khách hàng", "Thêm giỏ hàng và đặt hàng", "Đơn hàng được tạo, chi tiết đơn hàng được lưu và event mua hàng được ghi nhận."),
        ("UC07", "Khách hàng", "Đánh giá sản phẩm", "Review được lưu và có thể dùng cho phân tích chất lượng sản phẩm."),
        ("UC08", "Admin", "Quản lý catalog", "Admin thêm, sửa, ẩn hoặc cập nhật sản phẩm, danh mục, thương hiệu."),
        ("UC09", "Admin", "Xem dashboard", "Admin xem số liệu tổng hợp về doanh thu, hành vi, gợi ý và chatbot."),
        ("UC10", "Hệ thống", "Ghi nhận behavior event", "Event được lưu theo user/session/product/time để phục vụ phân tích."),
        ("UC11", "Hệ thống", "Cập nhật user profile", "Hồ sơ nhu cầu người dùng được cập nhật từ quiz và hành vi."),
        ("UC12", "Hệ thống", "Tính recommendation", "Điểm KB, CB, CF được kết hợp thành điểm hybrid cuối cùng."),
    ], widths=[0.65, 1.15, 2.15, 3.0])

    heading(doc, "7 Thiết kế cơ sở dữ liệu ban đầu", 1)
    para(doc, "Cơ sở dữ liệu ban đầu dùng MySQL. Thiết kế tập trung vào hai nhóm dữ liệu: dữ liệu thương mại điện tử và dữ liệu phục vụ hệ thống gợi ý, chatbot, analytics.")
    add_table(doc, ["Nhóm", "Bảng", "Mục đích"], [
        ("Ecommerce", "users", "Lưu tài khoản khách hàng và quản trị viên."),
        ("Ecommerce", "brands", "Lưu thương hiệu mỹ phẩm."),
        ("Ecommerce", "categories", "Lưu danh mục sản phẩm, hỗ trợ danh mục cha con."),
        ("Ecommerce", "products", "Lưu thông tin sản phẩm và thuộc tính dùng cho AI."),
        ("Ecommerce", "orders", "Lưu đơn hàng của khách hàng."),
        ("Ecommerce", "order_items", "Lưu từng sản phẩm trong đơn hàng."),
        ("AI/Data", "user_behavior_events", "Lưu hành vi người dùng như xem sản phẩm, tìm kiếm, lọc, thêm giỏ, mua hàng."),
        ("AI/Data", "user_profiles", "Lưu hồ sơ sở thích và nhu cầu chăm sóc da của người dùng."),
        ("AI/Data", "recommendation_logs", "Lưu request và kết quả gợi ý để đánh giá và debug."),
        ("AI/Data", "chatbot_conversations", "Lưu hội thoại chatbot và sản phẩm được đề xuất."),
    ], widths=[1.2, 2.0, 3.7])
    para(doc, "Migration ban đầu đã được đặt tại database/migrations/001_initial_schema.sql. File này tạo database cosmetic_ecommerce và các bảng chính cùng khóa ngoại, index phục vụ truy vấn hành vi theo user, product, event type và thời gian.")

    heading(doc, "8 Thiết kế API và event contract", 1)
    para(doc, "Để website và hệ thống bên trong có thể phát triển độc lập, Giai đoạn 1 đã tạo contract giao tiếp trước khi triển khai chi tiết. Website chỉ cần tuân thủ API và event schema đã thống nhất, còn các service hệ thống có thể thay đổi nội bộ mà không làm vỡ giao diện tích hợp.")
    add_table(doc, ["API hoặc Event", "Mục đích", "Phía sử dụng"], [
        ("POST /behavior-events", "Ghi nhận hành vi người dùng từ website.", "Website gửi, analytics/recommendation nhận dữ liệu."),
        ("POST /recommendations/personalized", "Lấy danh sách sản phẩm gợi ý cá nhân hóa.", "Website gọi recommendation-service."),
        ("GET /recommendations/similar-products/{productId}", "Lấy sản phẩm tương tự sản phẩm đang xem.", "Trang chi tiết sản phẩm."),
        ("POST /chatbot/messages", "Gửi tin nhắn người dùng và nhận phản hồi chatbot.", "Chatbot UI trên website."),
        ("GET /analytics/dashboard-summary", "Lấy số liệu tổng quan cho dashboard admin.", "Admin dashboard."),
    ], widths=[2.4, 2.55, 1.95])

    heading(doc, "9 Chuẩn bị Product Dataset", 1)
    heading(doc, "9.1 Thông tin cần thu thập", 2)
    add_table(doc, ["Nhóm dữ liệu", "Trường dữ liệu", "Mục đích"], [
        ("Định danh", "sku, name, brand, category", "Nhận diện sản phẩm và liên kết với catalog website."),
        ("Bán hàng", "price, volume, image_url", "Hiển thị trên website và hỗ trợ lọc theo ngân sách."),
        ("Mô tả", "description, benefits, usage_instruction", "Hiển thị chi tiết sản phẩm và hỗ trợ chatbot."),
        ("Thành phần", "inci_ingredients, key_ingredients", "Phân tích phù hợp với loại da, vấn đề da và thành phần cần tránh."),
        ("Cá nhân hóa", "skin_types, skin_concerns, care_goals", "Dữ liệu lõi cho Knowledge-Based và Content-Based Recommendation."),
        ("An toàn", "warnings", "Giúp chatbot trả lời có kiểm soát và cảnh báo khi cần."),
        ("Kiểm chứng", "source_url, verified_at", "Theo dõi nguồn dữ liệu và thời điểm xác minh."),
    ], widths=[1.35, 2.25, 3.3])

    heading(doc, "9.2 Phương án làm sạch và chuẩn hóa", 2)
    add_manual_numbered(doc, [
        "Lưu dữ liệu ban đầu vào datasets/raw để giữ nguyên nguồn thu thập.",
        "Chuẩn hóa tên thương hiệu, danh mục và mã SKU.",
        "Chuẩn hóa giá về số nguyên hoặc decimal theo đơn vị VND.",
        "Chuẩn hóa dung tích về dạng thống nhất như 50ml, 100ml, 150ml.",
        "Tách danh sách thành phần bằng dấu phân tách thống nhất.",
        "Map loại da, vấn đề da và mục tiêu chăm sóc về bộ enum nội bộ.",
        "Loại bỏ sản phẩm trùng SKU, trùng URL hoặc trùng tên trong cùng thương hiệu.",
        "Kiểm tra thiếu dữ liệu ở các trường quan trọng như brand, category, price, skin_types, skin_concerns, source_url.",
        "Ghi dữ liệu đã xử lý vào datasets/processed để phục vụ import database và mô hình gợi ý.",
    ])

    heading(doc, "10 Kết quả đã hoàn thành và kế hoạch tiếp theo", 1)
    add_table(doc, ["Hạng mục", "Trạng thái", "Vị trí trong source"], [
        ("Clone repository DACS về máy local", "Hoàn thành", "DACS"),
        ("Tách source theo monorepo", "Hoàn thành", "apps, services, packages, contracts, database, datasets, docs"),
        ("Tài liệu phân công và kiến trúc source", "Hoàn thành", "docs/architecture/source-split-and-ownership.md"),
        ("Kế hoạch Giai đoạn 1", "Hoàn thành", "docs/phase-1/phase-1-implementation-plan.md"),
        ("API contract", "Hoàn thành", "contracts/openapi/system-api.yaml"),
        ("Behavior event contract", "Hoàn thành", "contracts/events/behavior-events.md"),
        ("Database migration ban đầu", "Hoàn thành", "database/migrations/001_initial_schema.sql"),
        ("Data dictionary", "Hoàn thành", "database/docs/data-dictionary.md"),
        ("Product dataset schema", "Hoàn thành", "datasets/product-catalog/product_dataset_schema.csv"),
        ("Push GitHub", "Chưa thực hiện", "Chỉ push sau khi nhóm kiểm tra ổn định"),
    ], widths=[2.3, 1.25, 3.35])
    para(doc, "Sau Giai đoạn 1, nhóm có thể chuyển sang Giai đoạn 2 theo hai nhánh song song. Nhánh website bắt đầu dựng giao diện, routing và gọi API mock theo contract. Nhánh hệ thống bắt đầu triển khai database, seed data, behavior tracking, recommendation-service, chatbot-service và analytics-service.")
    para(doc, "Điểm quan trọng nhất của Giai đoạn 1 là ranh giới trách nhiệm đã rõ: website phụ trách trải nghiệm người dùng, còn hệ thống bên trong phụ trách dữ liệu, gợi ý, chatbot, phân tích và dashboard. Cách chia này giúp mỗi người làm phần của mình nhưng vẫn đảm bảo có điểm nối để tích hợp cuối cùng.")

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
