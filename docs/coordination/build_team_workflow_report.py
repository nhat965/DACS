from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "Tai_lieu_phan_cong_va_phoi_hop_du_an.docx"

BLACK = "000000"
DARK = "1F2933"
GRAY = "4B5563"
BLUE = "1F4E79"
LIGHT_BLUE = "EAF3F8"
BORDER = "D9D9D9"
HEADER_GRAY = "404040"


def set_font(run, size=10.2, bold=False, italic=False, color=DARK, name="Aptos"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def remove_style_borders(style):
    p_pr = style._element.find(qn("w:pPr"))
    if p_pr is None:
        return
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_cell_text(cell, text, bold=False, fill=None, color=DARK, size=8.8, align=None):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)
    if fill:
        shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align is None:
        align = WD_ALIGN_PARAGRAPH.CENTER if len(str(text)) <= 20 and "\n" not in str(text) else WD_ALIGN_PARAGRAPH.LEFT
    p.alignment = align
    r = p.add_run(str(text))
    set_font(r, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths=None, header_fill=BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for i, value in enumerate(headers):
        set_cell_text(header.cells[i], value, bold=True, fill=header_fill, color="FFFFFF", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            header.cells[i].width = Inches(widths[i])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        fill = "FFFFFF" if row_index % 2 == 0 else LIGHT_BLUE
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, fill=fill, size=8.55)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def para(doc, text="", lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    if lead:
        r1 = p.add_run(lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text)
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    remove_paragraph_borders(p)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        set_font(run, size=14.5 if level == 1 else 11.5 if level == 2 else 10.2, bold=True, color=BLACK, name="Aptos Display")
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(item)
        set_font(r)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(f"{index}.  {item}")
        set_font(r)


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_border(cell)
    set_cell_margins(cell, 120, 140, 120, 140)
    shade(cell, "F3F4F6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_font(r, size=8.4, name="Consolas", color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_font(run, size=8.2, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1

    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        remove_style_borders(style)
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.color.rgb = RGBColor.from_string(BLACK)

    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.2)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.14)

    header = section.header.paragraphs[0]
    header.text = "Tài liệu phân công và phối hợp dự án | DACS"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_font(run, size=8.2, color=GRAY)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(46)
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("Tài liệu phân công và phối hợp dự án")
    set_font(r, size=21, bold=True, color=BLACK, name="Aptos Display")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("Website thương mại điện tử mỹ phẩm tích hợp hệ thống gợi ý chatbot analytics và dashboard")
    set_font(r, size=12.5, bold=True, color=DARK, name="Aptos Display")

    add_table(doc, ["Thông tin", "Nội dung"], [
        ("Mục đích", "Làm tài liệu thống nhất cách chia việc, quy ước tên gọi, workflow tích hợp và checklist phối hợp giữa hai thành viên."),
        ("Phạm vi", "Website, backend/API, database, product dataset, recommendation, chatbot, analytics, dashboard và stream analytics."),
        ("Repository", "DACS"),
        ("Người 1", "Phụ trách website, giao diện khách hàng, giao diện admin và tích hợp API ở frontend."),
        ("Người 2", "Phụ trách database, dataset, API, recommendation, chatbot, analytics và dữ liệu dashboard."),
        ("Ngày lập", "12/09/2026"),
    ], widths=[1.5, 5.6], header_fill=HEADER_GRAY)

    para(doc, "Tài liệu này được dùng như bản thống nhất làm việc cho hai thành viên trong nhóm. Mục tiêu là mỗi người có thể phát triển phần của mình độc lập, nhưng vẫn dùng đúng tên field, đúng endpoint, đúng event và đúng format dữ liệu để khi ghép hệ thống không phải sửa lại nhiều.")

    doc.add_page_break()

    heading(doc, "1 Nguyên tắc phối hợp chung", 1)
    add_numbered(doc, [
        "Website và hệ thống bên trong phát triển tách source, nhưng phải tuân thủ cùng API contract và event contract.",
        "Người làm website được phép dùng mock data, nhưng mock data phải dùng đúng tên field, enum và cấu trúc response như API thật.",
        "Người làm hệ thống khi đổi database, API hoặc event phải cập nhật tài liệu contract trước khi yêu cầu frontend tích hợp.",
        "Không tự đặt tên field riêng trong frontend nếu field đó đã có trong contract.",
        "Chưa push GitHub hoặc merge phần lớn khi hai bên chưa test local tối thiểu các luồng chính.",
        "Các giá trị hiển thị tiếng Việt trên UI phải được map từ code tiếng Anh không dấu trong dữ liệu.",
    ])

    heading(doc, "2 Phân công nhiệm vụ", 1)
    add_table(doc, ["Người", "Khu vực source", "Trách nhiệm chính"], [
        ("Người 1", "apps/web", "Xây giao diện khách hàng, giao diện admin, routing, form, trạng thái loading/error/empty, gọi API thật hoặc mock theo contract."),
        ("Người 2", "database", "Thiết kế schema, migration, seed data, data dictionary và quy tắc đặt tên trong database."),
        ("Người 2", "datasets", "Thu thập, làm sạch, chuẩn hóa product dataset và import dữ liệu vào database."),
        ("Người 2", "services/recommendation-service", "Triển khai Knowledge-Based, Content-Based, Item-Based CF và Hybrid Recommendation."),
        ("Người 2", "services/chatbot-service", "Triển khai chatbot tư vấn dựa trên dữ liệu sản phẩm và quy tắc an toàn."),
        ("Người 2", "services/analytics-service", "Tổng hợp dữ liệu cho dashboard, hiệu quả gợi ý, hành vi người dùng và chatbot."),
        ("Người 2", "services/stream-analytics-service", "Thiết kế hoặc demo xử lý event gần thời gian thực nếu còn thời gian."),
        ("Cả hai", "contracts", "Thống nhất endpoint, request, response, enum, event type và error format."),
        ("Cả hai", "docs", "Cập nhật tài liệu phân tích, thiết kế, workflow, checklist và thay đổi quan trọng."),
    ], widths=[0.9, 2.1, 4.1])

    heading(doc, "3 Việc Người 1 cần làm", 1)
    add_table(doc, ["Nhóm việc", "Công việc cụ thể", "Đầu ra cần bàn giao"], [
        ("Khung website", "Chọn framework, tạo layout, header, footer, navigation và routing.", "Có app chạy local, điều hướng được các trang chính."),
        ("Khách hàng", "Làm trang chủ, danh mục, chi tiết sản phẩm, tìm kiếm, bộ lọc, giỏ hàng, checkout, đăng nhập, đăng ký, tài khoản.", "Các trang dùng mock data đúng contract và có responsive cơ bản."),
        ("Admin", "Làm dashboard, quản lý sản phẩm, danh mục, thương hiệu, đơn hàng, người dùng và đánh giá.", "Admin UI có bảng dữ liệu, form thêm/sửa và trạng thái rỗng/lỗi."),
        ("Tích hợp API", "Gọi product API, recommendation API, chatbot API, analytics API và gửi behavior event.", "API client dùng biến môi trường API base URL."),
        ("Behavior tracking", "Gửi event khi người dùng xem sản phẩm, tìm kiếm, lọc, click gợi ý, thêm giỏ, đặt hàng và đánh giá.", "Event body đúng `BehaviorEventRequest`."),
        ("UX hoàn thiện", "Làm loading state, empty state, error state, validate form và thông báo thao tác.", "Người dùng thao tác được flow demo trọn vẹn."),
    ], widths=[1.4, 3.3, 2.4])

    heading(doc, "4 Việc Người 2 cần làm", 1)
    add_table(doc, ["Nhóm việc", "Công việc cụ thể", "Đầu ra cần bàn giao"], [
        ("Database", "Hoàn thiện bảng users, brands, categories, products, orders, order_items, user_behavior_events, user_profiles, recommendation_logs, chatbot_conversations.", "Migration SQL chạy được và data dictionary cập nhật."),
        ("Dataset", "Thu thập dữ liệu mỹ phẩm, chuẩn hóa giá, dung tích, loại da, vấn đề da, mục tiêu chăm sóc, thành phần và nguồn dữ liệu.", "File raw, processed và script/import data nếu có."),
        ("API contract", "Bổ sung Product API, Auth API, Cart API, Order API, Admin API nếu cần cho website.", "OpenAPI cập nhật trước khi frontend tích hợp."),
        ("Behavior tracking", "Tạo endpoint nhận event và lưu vào database.", "POST /behavior-events trả 202 khi nhận event hợp lệ."),
        ("Recommendation", "Tính gợi ý cá nhân hóa, sản phẩm tương tự, lý do gợi ý và log kết quả.", "API recommendation trả đúng response contract."),
        ("Chatbot", "Nhận message, phân tích nhu cầu, truy vấn sản phẩm, trả lời an toàn và lưu hội thoại.", "API chatbot hoạt động với dữ liệu sản phẩm thật hoặc seed data."),
        ("Analytics", "Tính doanh thu, đơn hàng, sản phẩm top, từ khóa top, click recommendation và chỉ số chatbot.", "API dashboard trả dữ liệu cho admin UI."),
        ("Deploy backend", "Chuẩn bị cấu hình local và môi trường deploy backend nếu cần.", "Có API base URL để website gọi khi deploy Vercel."),
    ], widths=[1.4, 3.4, 2.3])

    heading(doc, "5 Việc cần làm chung", 1)
    add_table(doc, ["Hạng mục", "Hai người cần thống nhất", "Khi nào làm"], [
        ("Contract", "Endpoint, method, request body, response body, error format, enum và ví dụ mock.", "Trước khi frontend gọi API hoặc backend đổi response."),
        ("Database", "Tên bảng, tên cột, enum, quan hệ chính và dữ liệu seed dùng cho demo.", "Trước khi import dữ liệu và viết API."),
        ("Mock data", "Mock phải giống API thật về field name, kiểu dữ liệu và cấu trúc phân trang.", "Trước khi Người 1 dựng UI bằng mock."),
        ("Behavior event", "Event type, field bắt buộc, metadata theo từng event và thời điểm gửi.", "Trước khi tracking hành vi."),
        ("Luồng tích hợp", "Thứ tự tích hợp Product API, Auth, Cart, Order, Recommendation, Chatbot, Analytics.", "Trước tuần tích hợp."),
        ("Deploy", "Frontend deploy ở đâu, backend deploy ở đâu, biến môi trường API base URL.", "Trước khi demo online."),
        ("Demo", "Luồng demo, tài khoản admin, tài khoản khách hàng, sản phẩm mẫu, dữ liệu dashboard.", "Trước buổi báo cáo hoặc nghiệm thu."),
    ], widths=[1.45, 3.8, 1.85])

    heading(doc, "6 Cấu trúc source thống nhất", 1)
    add_code_block(doc, """DACS/
  apps/web
  services/recommendation-service
  services/chatbot-service
  services/analytics-service
  services/stream-analytics-service
  packages/shared-contracts
  contracts/openapi
  contracts/events
  database/migrations
  database/seeds
  database/docs
  datasets/raw
  datasets/processed
  datasets/product-catalog
  docs""")
    para(doc, "Người 1 chỉ cần làm chính trong `apps/web`, nhưng phải đọc `contracts` để mock và gọi API đúng. Người 2 làm chính trong `services`, `database`, `datasets` và `contracts`, nhưng phải giữ response ổn định để không làm vỡ giao diện.")

    heading(doc, "7 Quy ước đặt tên bắt buộc", 1)
    add_table(doc, ["Loại tên", "Quy ước", "Ví dụ đúng"], [
        ("Database table", "snake_case, số nhiều", "user_behavior_events, recommendation_logs"),
        ("Database column", "snake_case", "product_id, created_at, skin_type"),
        ("API JSON field", "camelCase", "productId, createdAt, skinType"),
        ("Endpoint", "plural noun, kebab-case nếu nhiều từ", "/behavior-events, /dashboard-summary"),
        ("Event type", "snake_case", "view_product, add_to_cart"),
        ("System enum", "UPPER_CASE", "CUSTOMER, ADMIN, ACTIVE, PENDING"),
        ("Domain code", "tiếng Anh không dấu, snake_case", "anti_acne, dark_spot"),
        ("UI label", "tiếng Việt", "Da dầu, Mụn, Chống nắng"),
    ], widths=[1.6, 2.6, 2.9])

    heading(doc, "8 Database cần thống nhất", 1)
    add_table(doc, ["Bảng", "Field chính", "Ghi chú phối hợp"], [
        ("users", "id, full_name, email, password_hash, role, skin_type, created_at, updated_at", "API không trả password_hash."),
        ("brands", "id, name, country, official_url, created_at", "Frontend dùng brandId và brandName khi hiển thị product."),
        ("categories", "id, name, parent_id, created_at", "Có thể bổ sung code nếu cần lọc ổn định hơn name."),
        ("products", "id, sku, name, brand_id, category_id, price, volume, stock_quantity, description, benefits, inci_ingredients, key_ingredients, skin_types, skin_concerns, care_goals, texture, usage_instruction, warnings, image_url, source_url, verified_at, status", "Đây là bảng quan trọng nhất cho website, recommendation và chatbot."),
        ("orders", "id, user_id, status, total_amount, created_at, updated_at", "Status dùng enum thống nhất."),
        ("order_items", "id, order_id, product_id, quantity, unit_price", "Dùng cho doanh thu và CF."),
        ("user_behavior_events", "id, user_id, session_id, product_id, event_type, event_value, metadata, occurred_at, created_at", "Frontend gửi event bằng camelCase, backend lưu snake_case."),
        ("user_profiles", "user_id, skin_type, skin_concerns, care_goals, preferred_categories, preferred_brands, avoid_ingredients, budget_min, budget_max, profile_confidence, updated_at", "Dùng cho recommendation và chatbot context."),
        ("recommendation_logs", "id, user_id, session_id, algorithm, request_context, result_items, created_at", "Dùng để đánh giá và debug gợi ý."),
        ("chatbot_conversations", "id, user_id, session_id, user_message, bot_reply, suggested_product_ids, created_at", "Dùng để phân tích chatbot."),
    ], widths=[1.55, 3.8, 1.75])

    heading(doc, "9 API endpoint cần thống nhất", 1)
    add_table(doc, ["Nhóm", "Endpoint", "Người dùng chính"], [
        ("Behavior", "POST /behavior-events", "Website gửi event, hệ thống lưu và phân tích."),
        ("Recommendation", "POST /recommendations/personalized", "Website lấy gợi ý cá nhân hóa."),
        ("Recommendation", "GET /recommendations/similar-products/{productId}", "Website lấy sản phẩm tương tự."),
        ("Chatbot", "POST /chatbot/messages", "Website gửi message và nhận reply."),
        ("Analytics", "GET /analytics/dashboard-summary", "Admin dashboard lấy số liệu tổng quan."),
        ("Product", "GET /products", "Website lấy danh sách sản phẩm."),
        ("Product", "GET /products/{productId}", "Website lấy chi tiết sản phẩm."),
        ("Catalog", "GET /brands, GET /categories", "Website lấy dữ liệu bộ lọc."),
        ("Auth", "POST /auth/register, POST /auth/login, GET /users/me, PUT /users/me", "Đăng ký, đăng nhập, tài khoản."),
        ("Cart", "GET /cart, POST /cart/items, PUT /cart/items/{cartItemId}, DELETE /cart/items/{cartItemId}", "Giỏ hàng."),
        ("Order", "POST /orders, GET /orders, GET /orders/{orderId}", "Đặt hàng và lịch sử đơn."),
        ("Review", "POST /products/{productId}/reviews", "Đánh giá sản phẩm."),
        ("Admin", "GET /admin/products, POST /admin/products, PUT /admin/products/{productId}, DELETE /admin/products/{productId}", "Quản lý sản phẩm."),
        ("Admin", "GET /admin/orders, PUT /admin/orders/{orderId}/status, GET /admin/users, GET /admin/dashboard", "Quản trị hệ thống."),
    ], widths=[1.2, 3.8, 2.1])

    heading(doc, "10 Field API và mock data cần dùng đúng", 1)
    heading(doc, "10.1 Product response", 2)
    add_code_block(doc, """{
  "id": 1,
  "sku": "COS-SAMPLE-001",
  "name": "Sample Hydrating Cleanser",
  "brandId": 1,
  "brandName": "Sample Brand",
  "categoryId": 1,
  "categoryName": "Cleanser",
  "price": 199000,
  "volume": "150ml",
  "stockQuantity": 20,
  "description": "Sua rua mat diu nhe",
  "benefits": "Lam sach, duong am",
  "inciIngredients": "water;glycerin;panthenol",
  "keyIngredients": "glycerin:duong_am|panthenol:lam_diu",
  "skinTypes": ["dry", "normal", "sensitive"],
  "skinConcerns": ["dryness", "redness"],
  "careGoals": ["hydrate", "repair"],
  "texture": "gel",
  "usageInstruction": "Dung sang va toi",
  "warnings": "Ngung dung khi kich ung",
  "imageUrl": "https://example.com/image.jpg",
  "status": "ACTIVE"
}""")
    heading(doc, "10.2 Product list response", 2)
    add_code_block(doc, """{
  "items": [],
  "page": 0,
  "size": 12,
  "totalItems": 100,
  "totalPages": 9
}""")
    heading(doc, "10.3 Behavior event request", 2)
    add_code_block(doc, """{
  "userId": 1,
  "sessionId": "session-abc",
  "productId": 10,
  "eventType": "view_product",
  "eventValue": 1,
  "metadata": {},
  "occurredAt": "2026-09-12T10:00:00+07:00"
}""")
    heading(doc, "10.4 Recommendation response", 2)
    add_code_block(doc, """{
  "algorithm": "hybrid",
  "items": [
    {
      "productId": 1,
      "score": 0.92,
      "reasons": ["Phu hop voi da dau", "Ho tro giam mun"]
    }
  ]
}""")
    heading(doc, "10.5 Chatbot response", 2)
    add_code_block(doc, """{
  "reply": "Ban co the tham khao cac san pham lam sach diu nhe phu hop voi da dau mun.",
  "suggestedProductIds": [1, 2, 3],
  "safetyNote": "Thong tin chi mang tinh tham khao, khong thay the tu van y te."
}""")

    heading(doc, "11 Enum và code dùng chung", 1)
    add_table(doc, ["Nhóm", "Giá trị thống nhất"], [
        ("User role", "CUSTOMER, ADMIN"),
        ("Product status", "ACTIVE, INACTIVE, DRAFT"),
        ("Order status", "PENDING, CONFIRMED, SHIPPING, COMPLETED, CANCELED"),
        ("Algorithm", "knowledge_based, content_based, item_based_cf, hybrid"),
        ("Skin type", "normal, dry, oily, combination, sensitive"),
        ("Skin concern", "acne, dark_spot, dryness, aging, redness, large_pores, dullness, uneven_texture, oiliness, sensitivity"),
        ("Care goal", "cleanse, hydrate, brighten, anti_acne, anti_aging, repair, soothe, oil_control, sun_protection, exfoliate"),
        ("Category code", "cleanser, toner, serum, moisturizer, sunscreen, exfoliant, mask, makeup_remover, eye_care, lip_care, body_care"),
        ("Sort", "newest, price_asc, price_desc, popular, best_selling, top_rated"),
    ], widths=[1.55, 5.55])

    heading(doc, "12 Event type và thời điểm gửi", 1)
    add_table(doc, ["Event type", "Website gửi khi nào", "Dùng cho hệ thống"], [
        ("view_product", "Người dùng mở trang chi tiết sản phẩm.", "Content-Based và analytics."),
        ("search_keyword", "Người dùng nhập và gửi tìm kiếm.", "User profile và từ khóa phổ biến."),
        ("filter_used", "Người dùng dùng bộ lọc.", "Knowledge-Based và analytics."),
        ("click_recommendation", "Người dùng click sản phẩm được gợi ý.", "Đánh giá hiệu quả recommendation."),
        ("add_favorite", "Người dùng thêm yêu thích.", "Content-Based."),
        ("add_to_cart", "Người dùng thêm giỏ hàng.", "Recommendation và conversion."),
        ("remove_from_cart", "Người dùng xóa khỏi giỏ.", "Analytics."),
        ("place_order", "Người dùng đặt hàng thành công.", "Collaborative Filtering và doanh thu."),
        ("review_product", "Người dùng đánh giá sản phẩm.", "Product quality và CF."),
        ("chatbot_message", "Người dùng gửi tin nhắn chatbot.", "Chatbot analytics."),
    ], widths=[1.55, 2.95, 2.6])

    heading(doc, "13 Query params cho danh sách sản phẩm", 1)
    add_table(doc, ["Param", "Ý nghĩa", "Ví dụ"], [
        ("keyword", "Từ khóa tìm kiếm.", "cleanser"),
        ("categoryId", "Lọc theo danh mục.", "1"),
        ("brandId", "Lọc theo thương hiệu.", "2"),
        ("minPrice", "Giá thấp nhất.", "100000"),
        ("maxPrice", "Giá cao nhất.", "500000"),
        ("skinType", "Lọc theo loại da.", "oily"),
        ("skinConcern", "Lọc theo vấn đề da.", "acne"),
        ("careGoal", "Lọc theo mục tiêu chăm sóc.", "hydrate"),
        ("sort", "Sắp xếp.", "price_asc"),
        ("page", "Trang hiện tại, bắt đầu từ 0.", "0"),
        ("size", "Số item mỗi trang.", "12"),
    ], widths=[1.55, 3.7, 1.85])
    add_code_block(doc, "GET /products?keyword=cleanser&skinType=oily&skinConcern=acne&page=0&size=12")

    heading(doc, "14 Error response chuẩn", 1)
    para(doc, "Mọi API nên trả lỗi cùng cấu trúc để frontend xử lý nhất quán.")
    add_code_block(doc, """{
  "timestamp": "2026-09-12T10:00:00+07:00",
  "status": 400,
  "error": "Bad Request",
  "message": "Invalid productId",
  "path": "/api/products/abc"
}""")

    heading(doc, "15 Workflow tích hợp giữa hai người", 1)
    add_table(doc, ["Bước", "Người 1 làm", "Người 2 làm", "Kết quả"], [
        ("1", "Đọc contract và dựng UI bằng mock data.", "Hoàn thiện contract và ví dụ request/response.", "Hai bên thống nhất tên field."),
        ("2", "Dựng trang sản phẩm, chi tiết và bộ lọc.", "Làm database, seed data, Product API.", "Website hiển thị được dữ liệu thật."),
        ("3", "Gửi behavior event từ UI.", "Làm endpoint /behavior-events.", "Hệ thống thu được hành vi người dùng."),
        ("4", "Tích hợp khối gợi ý.", "Làm recommendation API.", "Trang sản phẩm có gợi ý cá nhân hóa hoặc tương tự."),
        ("5", "Tích hợp chatbot UI.", "Làm chatbot API.", "Người dùng hỏi đáp sản phẩm được."),
        ("6", "Tích hợp admin dashboard.", "Làm analytics API.", "Admin xem số liệu tổng quan."),
        ("7", "Test flow end to end.", "Sửa lỗi API, dữ liệu và log.", "Có demo ổn định."),
    ], widths=[0.55, 2.25, 2.35, 1.95])

    heading(doc, "16 Workflow deploy", 1)
    add_numbered(doc, [
        "Frontend deploy lên Vercel hoặc môi trường hosting phù hợp cho website.",
        "Backend hoặc system services deploy ở môi trường riêng như Render, Railway, Fly.io, VPS hoặc chạy local khi demo.",
        "Database dùng MySQL local ở giai đoạn đầu, sau đó có thể chuyển lên dịch vụ cloud nếu cần demo online.",
        "Frontend không gọi vào thư mục contracts; frontend đọc theo contract rồi gọi API thật qua API base URL.",
        "Khi deploy frontend, cấu hình biến môi trường `VITE_API_BASE_URL` hoặc `NEXT_PUBLIC_API_BASE_URL`.",
        "Không để password, token, database URL thật trong repository.",
    ])
    add_code_block(doc, """Local:
VITE_API_BASE_URL=http://localhost:8080/api

Deploy:
VITE_API_BASE_URL=https://api-domain.com/api""")

    heading(doc, "17 Checklist trước khi bàn giao từng phần", 1)
    add_table(doc, ["Phần", "Checklist bắt buộc"], [
        ("Frontend page", "Có loading, empty, error state; dùng đúng mock field; responsive cơ bản; không hard-code API URL."),
        ("API endpoint", "Có request/response đúng contract; có status code đúng; có error response chuẩn; có ví dụ test."),
        ("Database migration", "Chạy được từ đầu; có khóa chính, khóa ngoại, index cần thiết; không mất dữ liệu seed quan trọng."),
        ("Dataset", "Có nguồn dữ liệu; có ngày xác minh; dữ liệu raw và processed tách riêng; không trùng SKU/URL."),
        ("Recommendation", "Trả productId, score, reasons; có log; xử lý user mới; không trả sản phẩm inactive."),
        ("Chatbot", "Không chẩn đoán bệnh; không cam kết điều trị; có safetyNote; câu trả lời dựa trên dữ liệu."),
        ("Analytics", "Số liệu có khoảng thời gian; query không quá nặng; response đúng dashboard contract."),
        ("Deploy", "Có env vars; CORS đúng domain frontend; API health check hoạt động; không lộ secrets."),
    ], widths=[1.55, 5.55])

    heading(doc, "18 Quy trình khi cần đổi field API hoặc database", 1)
    add_numbered(doc, [
        "Người muốn đổi field phải ghi rõ lý do đổi và phần bị ảnh hưởng.",
        "Cập nhật `contracts/openapi/system-api.yaml` hoặc migration/data dictionary trước.",
        "Báo cho người còn lại biết field cũ, field mới và thời điểm đổi.",
        "Nếu frontend đang dùng mock data, cập nhật mock data cùng lúc với contract.",
        "Nếu backend đã có dữ liệu, chuẩn bị migration hoặc mapping tương thích.",
        "Test lại endpoint liên quan trước khi tích hợp tiếp.",
    ])

    heading(doc, "19 Thứ tự ưu tiên triển khai", 1)
    add_table(doc, ["Giai đoạn", "Người 1", "Người 2", "Mục tiêu"], [
        ("Tuần 1", "Layout, trang sản phẩm, chi tiết, mock data.", "Database schema, seed data, contract đầy đủ.", "Có UI và dữ liệu mẫu thống nhất."),
        ("Tuần 2", "Giỏ hàng, checkout, auth UI, admin UI cơ bản.", "Product API, Auth API, Cart/Order API, behavior event API.", "Có luồng mua hàng cơ bản."),
        ("Tuần 3", "Recommendation UI, chatbot UI, dashboard UI.", "Recommendation, chatbot, analytics.", "Có phần AI và dashboard."),
        ("Tuần 4", "Tích hợp API thật, sửa lỗi UI.", "Sửa lỗi API, tối ưu dữ liệu, hỗ trợ deploy.", "Có demo end to end."),
    ], widths=[1.0, 2.1, 2.45, 1.55])

    heading(doc, "20 Luồng demo cuối cùng", 1)
    add_numbered(doc, [
        "Khách hàng mở website và xem danh sách sản phẩm.",
        "Khách hàng lọc theo loại da hoặc vấn đề da.",
        "Khách hàng xem chi tiết sản phẩm, hệ thống ghi event `view_product`.",
        "Website hiển thị sản phẩm tương tự hoặc gợi ý cá nhân hóa.",
        "Khách hàng thêm sản phẩm vào giỏ, hệ thống ghi event `add_to_cart`.",
        "Khách hàng đặt hàng, hệ thống ghi event `place_order`.",
        "Khách hàng hỏi chatbot về sản phẩm phù hợp.",
        "Admin mở dashboard để xem doanh thu, top sản phẩm và hiệu quả gợi ý.",
    ])

    heading(doc, "21 Tóm tắt để nói với bạn cùng nhóm", 1)
    para(doc, "Dự án sẽ làm theo hướng tách việc rõ ràng. Bạn phụ trách website trong `apps/web`; mình phụ trách hệ thống bên trong gồm database, dataset, API, recommendation, chatbot và analytics. Hai phần không phụ thuộc source trực tiếp của nhau, nhưng bắt buộc thống nhất qua `contracts`. Khi bạn làm mock data, mock phải giống API thật về tên field và cấu trúc response. Khi mình làm API thật, mình phải giữ đúng contract. Nếu cần đổi tên field hoặc endpoint, hai bên cập nhật contract trước rồi mới sửa code.")

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
